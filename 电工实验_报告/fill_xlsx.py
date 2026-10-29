# -*- coding: utf-8 -*-
"""在源数据基础上另存“含计算”工作簿（不改动源文件）。
- 实验二：补相对误差、计算值，新增“等效变换”表（按实验册第 3--4 页属实验二）
- 实验三：补 cosφ 计算列，新增分析表与电压三角形校验
"""
import json
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = r"D:\aaa_D_usual\DSH\111学习助手\电工实验_报告"
with open(os.path.join(ROOT, "计算结果.json"), encoding="utf-8") as fh:
    D = json.load(fh)

HEAD = PatternFill("solid", fgColor="DDEBF7")
CALC = PatternFill("solid", fgColor="FDE9D9")
NOTE = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
f2 = lambda v, n=2: round(float(v), n)


def style(ws, r, c, v=None, fmt="0.00", fill=None, bold=False):
    cell = ws.cell(row=r, column=c)
    if v is not None:
        cell.value = v
    cell.border = BOX
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.number_format = fmt
    if fill:
        cell.fill = fill
    if bold:
        cell.font = Font(bold=True)


# ==================== 实验二 ====================
wb = openpyxl.load_workbook(os.path.join(ROOT, "实验2_直流电路数据.xlsx"))

# 2.1.2 相对误差行
ws = wb["2.1.2 电阻测量"]
for j, v in enumerate(D["t212"]["error"], start=2):
    style(ws, 4, j, f2(v, 2), fill=CALC)

# 2.1.3 计算值 / 相对误差
ws = wb["2.1.3 KVL-KCL"]
calc = D["t213"]["calc"]
order = ["I1", "I2", "I3", "U_AB", "U_CD", "U_AD", "U_DE", "U_FA"]
style(ws, 3, 2, 6.0, fill=CALC)
style(ws, 3, 3, 12.0, fill=CALC)
for j, key in enumerate(order, start=4):
    style(ws, 3, j, f2(calc[key], 2), fill=CALC)
rel = D["t213"]["rel"]
style(ws, 4, 2, f2(abs(5.9 - 6) / 6 * 100, 1), fmt="0.0", fill=CALC)
style(ws, 4, 3, f2(abs(12.1 - 12) / 12 * 100, 1), fmt="0.0", fill=CALC)
for j, key in enumerate(order, start=4):
    style(ws, 4, j, f2(rel[key], 1), fmt="0.0", fill=CALC)

# 2.1.4 校验新表
if "2.1.4 校验" in wb.sheetnames:
    del wb["2.1.4 校验"]
ws = wb.create_sheet("2.1.4 校验")
cols = ["I1/mA", "I2/mA", "I3/mA", "U_AB/V", "U_CD/V", "U_AD/V", "U_DE/V", "U_FA/V"]
keys = ["I1", "I2", "I3", "U_AB", "U_CD", "U_AD", "U_DE", "U_FA"]
ws.cell(row=1, column=1, value="叠加定理校验（由 2.1.4 实测数据算出）").font = Font(bold=True, size=11)
hdr = ["项目"] + cols
for j, h in enumerate(hdr, start=1):
    style(ws, 2, j, h, fmt="General", fill=HEAD, bold=True)
style(ws, 3, 1, "可加性残差（行1+行2−行3）", fmt="General")
for j, key in enumerate(keys, start=2):
    style(ws, 3, j, f2(D["t214"]["sum_res"][key], 1), fmt="0.0", fill=CALC)
style(ws, 4, 1, "齐次性残差（2×行2−行4）", fmt="General")
for j, key in enumerate(keys, start=2):
    style(ws, 4, j, f2(D["t214"]["hom_res"][key], 1), fmt="0.0", fill=CALC)
style(ws, 6, 1, "理论值（用实测电阻正演，供对照）", fmt="General", bold=True)
for j, h in enumerate(["工况"] + cols, start=1):
    style(ws, 7, j, h, fmt="General", fill=HEAD, bold=True)
for i, (nm, key) in enumerate(zip(["U_S1 单独", "U_S2 单独", "共同作用", "2U_S2"], D["t214"]["cases"])):
    style(ws, 8 + i, 1, nm, fmt="General")
    th = D["t214"]["theory"][key]
    for j, kk in enumerate(keys, start=2):
        style(ws, 8 + i, j, f2(th[kk], 2), fill=CALC)
ws.cell(row=13, column=1,
        value="结论：除 U_AB 一列残差 0.5 V（源于第 1 行读数偏高）外，可加性与齐次性均成立；"
              "功率不能用叠加定理计算（P=I²R 为二次关系）。").font = Font(size=9, bold=True, color="C00000")
ws.column_dimensions["A"].width = 26
for j in range(2, 10):
    ws.column_dimensions[get_column_letter(j)].width = 10.5

# 等效变换（从实验三源文件读取，移入实验二）
if "等效变换" in wb.sheetnames:
    del wb["等效变换"]
src = openpyxl.load_workbook(os.path.join(ROOT, "实验3_交流电路数据.xlsx"))
ws3 = src["3-1 等效变换"]
eq = D["equiv"]
ws = wb.create_sheet("等效变换")
ws.cell(row=1, column=1, value="3. Verification of Equivalent-transformation Conditions（实验册第 3--4 页）").font = Font(bold=True, size=11)
hdr = ["电路", "电源取值", "端电压 U/V", "端电流 I/mA", "端口电阻 U/I /Ω", "说明"]
for j, h in enumerate(hdr, start=1):
    style(ws, 2, j, h, fmt="General", fill=HEAD, bold=True)
rows = [("图一（电压源 + R0 串联）", ws3.cell(row=2, column=2).value, ws3.cell(row=2, column=3).value,
         ws3.cell(row=2, column=4).value, f2(eq["R_port1"], 1), "R0=51.3Ω, RL=216.4Ω"),
        ("图二（电流源 + R0 并联）", ws3.cell(row=3, column=2).value, ws3.cell(row=3, column=3).value,
         ws3.cell(row=3, column=4).value, f2(eq["R_port2"], 1), "R0=51.3Ω, RL=216.4Ω")]
for i, row in enumerate(rows, start=3):
    for j, v in enumerate(row, start=1):
        style(ws, i, j, v, fmt=("0.0" if j >= 3 else "General"), fill=(CALC if j == 5 else None))
notes = [
    "判据：两种接法接同一负载时端电压、端电流相同 → 对外等效。实测均为 4.1 V / 19.2 mA，等效成立。",
    f"等效条件 I_S = U_S/R0：理论 I_S = 6.0/51.3 = {eq['IS_theory']:.1f} mA，实测 103.5 mA（差约 12%，来自电源内阻与表内阻）。",
    f"图二理论校验：I_L = I_S·R0/(R0+RL) = {eq['I_theory2']:.2f} mA（实测 19.2 mA），U_L = {eq['U_theory2']:.2f} V（实测 4.1 V）。",
]
for i, t in enumerate(notes, start=6):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(size=9)
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
for j, w in enumerate([26, 18, 12, 12, 14, 26], start=1):
    ws.column_dimensions[get_column_letter(j)].width = w

p2 = os.path.join(ROOT, "实验2_数据与计算.xlsx")
wb.save(p2)

# ==================== 实验三 ====================
wb = openpyxl.load_workbook(os.path.join(ROOT, "实验3_交流电路数据.xlsx"))
ws = wb["3-2 日光灯"]
for i, r in enumerate(D["fluor"]["rows"], start=2):
    style(ws, i, 11, f2(r["cos_calc"], 3), fmt="0.000", fill=CALC)

if "3-2 分析" in wb.sheetnames:
    del wb["3-2 分析"]
ws = wb.create_sheet("3-2 分析")
ws.cell(row=1, column=1, value="日光灯电路补充计算（S、Q、I_C 理论值、偏差）").font = Font(bold=True, size=11)
hdr = ["C/μF", "S=UI /VA", "Q=√(S²−P²) /var", "I_C理=ωCU /mA", "cosφ计算", "cosφ实测", "偏差/%"]
for j, h in enumerate(hdr, start=1):
    style(ws, 2, j, h, fmt="General", fill=HEAD, bold=True)
for i, r in enumerate(D["fluor"]["rows"], start=3):
    style(ws, i, 1, r["C"], fmt="0.0")
    style(ws, i, 2, f2(r["S"], 1), fmt="0.0", fill=CALC)
    style(ws, i, 3, f2(r["Q"], 1), fmt="0.0", fill=CALC)
    style(ws, i, 4, f2(r["IC_th"], 0), fmt="0", fill=CALC)
    style(ws, i, 5, f2(r["cos_calc"], 3), fmt="0.000")
    style(ws, i, 6, f2(r["cos_meas"], 3), fmt="0.000")
    style(ws, i, 7, f2(r["dev"], 1), fmt="0.0")
ws.cell(row=10, column=1,
        value=f"最佳补偿电容 C_opt = P·tanφ₀/(ωU²) = {D['fluor']['C_opt']:.2f} μF（φ₀=65.6°），"
              f"与实测最小电流对应的 C=3.7 μF 一致。").font = Font(size=9, bold=True, color="C00000")
ws.column_dimensions["A"].width = 9
for j in range(2, 8):
    ws.column_dimensions[get_column_letter(j)].width = 15

if "3-3 校验" in wb.sheetnames:
    del wb["3-3 校验"]
ws = wb.create_sheet("3-3 校验")
ws.cell(row=1, column=1, value="电压三角形校验：√(U_R²+U_C²) 与实测 U 比较").font = Font(bold=True, size=11)
hdr = ["", "C=1 μF", "C=2 μF", "C=4.7 μF"]
for j, h in enumerate(hdr, start=1):
    style(ws, 2, j, h, fmt="General", fill=HEAD, bold=True)
names = ["U_R/V", "U_C/V", "U实测/V", "√(U_R²+U_C²)/V", "偏差/%", "φ=arctan(U_C/U_R)/°"]
for j, nm in enumerate(names, start=1):
    style(ws, 3 + j - 1, 1, nm, fmt="General", bold=True)
for c, r in enumerate(D["vtri"]["rows"], start=2):
    style(ws, 3, c, f2(r["UR"], 1), fmt="0.0")
    style(ws, 4, c, f2(r["UC"], 1), fmt="0.0")
    style(ws, 5, c, f2(r["U"], 1), fmt="0.0")
    style(ws, 6, c, f2(r["U_calc"], 1), fmt="0.0", fill=CALC)
    style(ws, 7, c, f2(r["dev"], 2), fmt="0.00", fill=CALC)
    style(ws, 8, c, f2(r["phi"], 1), fmt="0.0", fill=CALC)
ws.column_dimensions["A"].width = 22
for j in range(2, 5):
    ws.column_dimensions[get_column_letter(j)].width = 12

# 说明：等效变换已移到实验二工作簿
ws = wb["说明"]
r = ws.max_row + 1
ws.cell(row=r, column=1, value="等效变换表的位置").font = Font(bold=True)
ws.cell(row=r, column=2,
        value="按实验册页码，等效变换条件验证在实验二第 3--4 页，已移入《实验2_数据与计算.xlsx》的“等效变换”表；"
              "本文件仍保留源表 3-1 以备查。")
ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[r].height = 44

p3 = os.path.join(ROOT, "实验3_数据与计算.xlsx")
wb.save(p3)
print("saved:", p2)
print("saved:", p3)
